"""Проверка инфраструктуры по внешним базам на синтетическом KSC (KAVTEST): сопоставление с каталогом, компьютеры из реестра ПО, отчёт XLSX/HTML, совпадение просмотра и отчёта.
Данные уязвимостей — искусственные записи (CVE-2099-*) в локальной базе. Без KSC_TEST_HOST пропускается."""
import uuid

import pytest
from openpyxl import load_workbook

from app.models import utcnow
from app.reports.runner import job_dir
from app.vdb import store
from app.vdb.normalize import Aff, Rec
from tests.test_audit_e2e import _run, _setup
from tests.test_ksc_live import HOST

pytestmark = pytest.mark.skipif(not HOST, reason="нет синтетической базы KSC (KSC_TEST_HOST)")
# Contoso Agent: 8.4.14 — компьютеры 1, 5 (без связи), 9 (кластер); 8.4.15 — 2, 3; 8.9 — 7; 8.10 — 6, 8. Google Chrome: 124.0.6367.60 — компьютер 6; 120.0.6099.130 — 7 и 8 (8 без связи).


async def seed(app):
    now = utcnow()
    async with app.state.session_maker() as db:
        await store.ensure_sources(db)
        nvd = [Rec(source="nvd", ext_id="CVE-2099-1", vuln_key="CVE-2099-1", severity="critical", score=9.8, title="Chrome: выполнение кода", published=now, modified=now,
                   affected=[Aff(vendor="google", product="chrome", v_end="121")]),
               Rec(source="nvd", ext_id="CVE-2099-2", vuln_key="CVE-2099-2", severity="high", score=8.0, title="Contoso Agent: повышение привилегий", published=now, modified=now, fix_available=True,
                   affected=[Aff(vendor="contoso", product="agent", v_end="8.4.15")]),
               Rec(source="nvd", ext_id="CVE-2099-3", vuln_key="CVE-2099-3", severity="high", score=7.5, title="Другой продукт с похожим названием", published=now, modified=now,
                   affected=[Aff(vendor="contoso", product="agent_manager", versions_all=True)])]
        await store.save_batch(db, 1, nvd)
        await store.save_batch(db, 1, [Rec(source="kev", ext_id="CVE-2099-2", vuln_key="CVE-2099-2", exploited=True, published=now, modified=now, affected=[Aff(vendor="Contoso", product="Agent", versions_all=True)])])
        await store.refresh_vulns(db, {"CVE-2099-1", "CVE-2099-2", "CVE-2099-3"})
        await db.commit()


async def preview(client, params):
    r = await client.post("/api/v1/quick/preview", json={"kind": "quick_extvulns", "params": params})
    assert r.status_code == 200, r.text
    return r.json()


def metric(p, key):
    return next(m["value"] for m in p["doc"]["metrics"] if m["key"] == key)


async def test_infrastructure_check_hosts_statuses_and_no_false_positive(client, app):
    sid, runner = await _setup(client, app)
    await seed(app)
    p = await preview(client, {"vuln_keys": ["CVE-2099-1", "CVE-2099-2", "CVE-2099-3"]})
    hosts = {r["host_id"]: r for r in p["items"]}
    by_vuln = {}
    for r in p["items"]:
        by_vuln.setdefault(r["vuln_key"], set()).add(r["host_id"])
    assert by_vuln["CVE-2099-1"] == {7} and by_vuln["CVE-2099-2"] == {1} and "CVE-2099-3" not in by_vuln          # Chrome 124 не затронут, 5 и 8 — без связи, 9 — кластер, похожее название не совпало
    assert metric(p, "hosts") == 2 and metric(p, "vulns") == 2 and metric(p, "installs") == 2 and metric(p, "kev") == 1
    r1 = next(r for r in p["items"] if r["vuln_key"] == "CVE-2099-1")
    assert r1["status"] == "confirmed" and r1["kev"] == "" and r1["installed_version"].startswith("120.") and r1["severity_ru"] == "Критическая"
    r2 = next(r for r in p["items"] if r["vuln_key"] == "CVE-2099-2")
    assert r2["kev"] == "да" and r2["status"] == "confirmed" and r2["installed_version"] == "8.4.14"                  # KEV не расширяет версии, границы — от NVD
    assert p["items"][0]["vuln_key"] == "CVE-2099-2"                                                                      # эксплуатируемая — первой
    assert any("НЕ подтверждены" in n for n in p["doc"]["notes"]) and any("кластеров" in n for n in p["doc"]["notes"])
    allc = await preview(client, {"vuln_keys": ["CVE-2099-1"], "include_inactive": True})
    assert {r["host_id"] for r in allc["items"]} == {7, 8}
    allh = await preview(client, {"vuln_keys": ["CVE-2099-2"], "include_inactive": True})
    assert {r["host_id"] for r in allh["items"]} == {1, 5}                                                                # устройство без связи — только если включены все
    miss = await preview(client, {"vuln_keys": ["CVE-2099-2", "CVE-9999-0"]})
    assert any("CVE-9999-0" in n for n in miss["doc"]["notes"])


async def test_by_product_selection_and_filters(client, app):
    sid, runner = await _setup(client, app)
    await seed(app)
    p = await preview(client, {"product_keys": ["contoso-agent"]})
    assert {r["vuln_key"] for r in p["items"]} == {"CVE-2099-2"} and metric(p, "hosts") == 1                              # все известные уязвимости продукта
    assert (await preview(client, {"product_keys": ["google-chrome"], "min_severity": "critical"}))["total"] == 1
    assert (await preview(client, {"product_keys": ["google-chrome"], "exploited_only": True}))["total"] == 0
    assert (await preview(client, {"vuln_keys": ["CVE-2099-2"], "statuses": ["fixed"]}))["total"] == 0
    bad = await client.post("/api/v1/quick/preview", json={"kind": "quick_extvulns", "params": {}})
    assert bad.status_code == 422


async def test_report_files_match_preview(client, app):
    sid, runner = await _setup(client, app)
    await seed(app)
    params = {"vuln_keys": ["CVE-2099-1", "CVE-2099-2"]}
    pv = await preview(client, params)
    job = (await client.post("/api/v1/reports", json={"kind": "quick_extvulns", "title": "", "params": params, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    done = (await client.get(f"/api/v1/reports/{job['id']}")).json()
    assert done["status"] == "done" and done["rows"] == pv["total"] and [m["value"] for m in done["summary"]["metrics"]] == [m["value"] for m in pv["doc"]["metrics"]]
    d = job_dir(app.state.settings, uuid.UUID(job["id"]))
    wb = load_workbook(d / "report.xlsx", read_only=True)
    assert wb.sheetnames == ["Сводка", "Компьютеры", "Уязвимости", "Программы", "Обнаружения", "Параметры отчёта"]
    rows = list(wb["Обнаружения"].iter_rows(values_only=True))
    assert rows[0][:7] == ("Опасность", "Уязвимость", "Идентификаторы", "Название", "Эксплуатация (CISA KEV)", "Результат сопоставления", "Основание")
    assert len(rows) - 1 == pv["total"] and {r[5] for r in rows[1:]} == {"Подтверждено источником"}
    html = (d / "report.html").read_text(encoding="utf-8")
    assert "НЕ подтверждены" in html and "Подтверждено источником" in html and "CVE-2099-2" in html and "<script" not in html
