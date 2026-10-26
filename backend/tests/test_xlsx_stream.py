"""XLSX пишется потоком: память не растёт с числом строк, содержимое листов и форматы сохраняются."""
import tracemalloc
from datetime import datetime, timezone

from openpyxl import load_workbook

from app.reports import builder, export


def _rows(n):
    return [{k: (f"v{i}-{k}" if "date" not in k and "connected" not in k else datetime(2026, 1, 1, tzinfo=timezone.utc)) for k, _ in builder.SOFTWARE_COLUMNS} | {"host_id": i % 500, "name": f"h{i % 500}", "linux_release": "", "branch": "7.4",
                                                                                                                                                                  "app_name": "php", "app_version": "7.4.1", "activity": "активен"} for i in range(n)]


def test_software_report_content_and_filters(tmp_path):
    rows = _rows(300)
    p = tmp_path / "r.xlsx"
    export.write_xlsx(p, "software_hosts", rows, {"fetched_at": "2026-01-01T00:00:00+00:00", "rows": 300, "hosts": 300, "branches": [{"branch": "7.4", "installs": 300, "hosts": 300, "packages": ["php"]}]},
                      {"Отчёт": "Проверка", "Источник KSC": "A", "_org": {"short_name": "Орг"}})
    wb = load_workbook(p)
    assert wb.sheetnames == ["Сводка", "Версии", "Компьютеры", "Установки", "Параметры отчёта"]
    ws = wb["Установки"]
    assert ws.max_row == 301 and ws.freeze_panes == "A2" and ws.auto_filter.ref == f"A1:{ws.cell(1, len(builder.SOFTWARE_COLUMNS)).column_letter}301"
    assert wb["Сводка"]["A1"].value == "Орг" and wb["Сводка"]["A1"].font.bold


def test_memory_does_not_grow_with_rows(tmp_path):
    rows = _rows(6000)
    tracemalloc.start()
    base = tracemalloc.get_traced_memory()[0]
    export.write_xlsx(tmp_path / "big.xlsx", "software_hosts", rows, {"fetched_at": "x", "rows": 6000, "hosts": 500, "branches": []}, {"Отчёт": "t", "_org": None})
    peak = tracemalloc.get_traced_memory()[1] - base
    tracemalloc.stop()
    assert peak < 25 * 1024 * 1024, f"пик {peak / 1048576:.0f} МБ: запись не потоковая"          # обычная книга openpyxl давала бы сотни мегабайт на таких данных
