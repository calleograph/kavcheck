"""Быстрые контекстные отчёты на синтетической базе KAVTEST с известными данными (tests/ksc_synthetic/schema.sql): уязвимости по CVE, по программам и по
компьютерам, подсчёт без двойного учёта, давность связи и «в сети сейчас», выгрузки XLSX/HTML/CSV и их совпадение с просмотром."""
import csv
import io
import re
import uuid

import pytest
from openpyxl import load_workbook

from app.models import ReportJob
from app.reports.runner import job_dir
from tests.test_audit_e2e import _run, _setup
from tests.test_ksc_live import HOST

pytestmark = pytest.mark.skipif(not HOST, reason="нет синтетической базы KSC (KSC_TEST_HOST)")
# KLA10001 (5001, CVE-2024-0001/0002): 7, 6, 1 активны; 8 и 5 без связи; 9 — кластер. KLA10903 (5903): 6 и 2. Chrome (vp 7001): 6 (KLA10001, KLA10002) и 7; 7 не в сети.


async def preview(client, kind, params):
    r = await client.post("/api/v1/quick/preview", json={"kind": kind, "params": params})
    assert r.status_code == 200, r.text
    return r.json()


def m(doc, key):
    return next(x["value"] for x in doc["doc"]["metrics"] if x["key"] == key)


async def test_vulns_by_cve_activity_and_online(client, app):
    await _setup(client, app)
    p = await preview(client, "quick_vulns", {"cves": ["CVE-2024-0001"]})
    assert m(p, "hosts") == 3 and m(p, "rows") == 3 and p["total"] == 3 and {r["host_id"] for r in p["items"]} == {1, 6, 7}
    assert all(r["cve"] == "CVE-2024-0001, CVE-2024-0002" and r["severity_ru"] == "Критическая" for r in p["items"])
    assert any("без связи" in n for n in p["doc"]["notes"]) and any("кластер" in n for n in p["doc"]["notes"])                  # скрытые и не-компьютеры объяснены
    assert {r["online"] for r in p["items"]} == {"в сети", "не в сети"} and next(r for r in p["items"] if r["host_id"] == 7)["active"] == "да"    # 7: связь за 30 суток, но не в сети
    allh = await preview(client, "quick_vulns", {"cves": ["CVE-2024-0001"], "include_inactive": True})
    assert m(allh, "hosts") == 5 and {r["host_id"] for r in allh["items"]} == {1, 5, 6, 7, 8}
    assert next(r for r in allh["items"] if r["host_id"] == 8)["online"] == "не в сети" and next(r for r in allh["items"] if r["host_id"] == 8)["active"] == "нет"
    on = await preview(client, "quick_vulns", {"cves": ["CVE-2024-0001"], "online_only": True})
    assert m(on, "hosts") == 2 and {r["host_id"] for r in on["items"]} == {1, 6} and any("не видит в сети" in n or "не в сети" in n for n in on["doc"]["notes"])
    assert m(on, "online") == 2


async def test_multi_selection_has_no_double_counting(client, app):
    await _setup(client, app)
    p = await preview(client, "quick_vulns", {"descr_ids": [5001, 5903]})        # KLA10001 на 7, 6, 1 и KLA10903 на 6, 2
    assert m(p, "hosts") == 4 and m(p, "rows") == 5 and m(p, "installs") == 5 and m(p, "descr") == 2          # компьютер 6 — один, хотя затронут обеими
    t = {t["id"]: t for t in p["doc"]["tables"]}
    assert t["hosts"]["n"] == 4 and t["vulns"]["n"] == 2 and t["programs"]["n"] == 2
    levels = {lv["key"]: lv["value"] for lv in p["doc"]["severity"]["levels"]}
    assert sum(levels.values()) == m(p, "hosts") and levels["critical"] == 3 and levels["warning"] == 1          # по наивысшему уровню: 7, 6, 1 — критическая; 2 — предупреждение
    # у каждой уязвимости — свои компьютеры, и их сумма больше уникальных (6 в обеих)
    vul = await client.post("/api/v1/reports", json={"kind": "quick_vulns", "title": "", "params": {"descr_ids": [5001, 5903]}, "source_ids": []})
    assert vul.status_code == 201


async def test_vulns_by_program_and_by_host(client, app):
    await _setup(client, app)
    by_prog = await preview(client, "quick_vulns", {"vp_ids": [7001]})
    assert m(by_prog, "hosts") == 2 and m(by_prog, "rows") == 3 and m(by_prog, "descr") == 2 and {r["host_id"] for r in by_prog["items"]} == {6, 7}
    assert {r["kla_id"] for r in by_prog["items"]} == {"KLA10001", "KLA10002"} and all(r["product"] == "Google Chrome" for r in by_prog["items"])
    host6 = await preview(client, "quick_vulns", {"host_ids": [6]})
    lv = {x["key"]: x["hint"] for x in host6["doc"]["severity"]["levels"]}
    assert m(host6, "hosts") == 1 and m(host6, "rows") >= m(host6, "installs") >= 2
    assert {r["severity"] for r in host6["items"]} >= {"critical", "high", "warning"} and lv
    only_high = await preview(client, "quick_vulns", {"host_ids": [6], "severity": "high"})
    assert {r["severity"] for r in only_high["items"]} == {"high"}
    empty = await preview(client, "quick_vulns", {"cves": ["CVE-0000-0000"]})
    assert empty["total"] == 0 and any("нет уязвимости" in n for n in empty["doc"]["notes"])
    bad = await client.post("/api/v1/quick/preview", json={"kind": "quick_vulns", "params": {"severity": "bogus"}})
    assert bad.status_code == 422                                              # без выбора отчёт строится по условиям страницы (см. тест ниже); неверные значения отклоняются


async def test_exports_match_preview_and_are_standalone(client, app):
    sid, runner = await _setup(client, app)
    params = {"descr_ids": [5001, 5903], "include_inactive": False}
    pv = await preview(client, "quick_vulns", params)
    job = (await client.post("/api/v1/reports", json={"kind": "quick_vulns", "title": "", "params": params, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    done = (await client.get(f"/api/v1/reports/{job['id']}")).json()
    assert done["status"] == "done" and done["rows"] == pv["total"] == 5 and done["summary"]["hosts"] == m(pv, "hosts") == 4
    assert [x["value"] for x in done["summary"]["metrics"]] == [x["value"] for x in pv["doc"]["metrics"]]                              # показатели просмотра и отчёта совпадают
    d = job_dir(app.state.settings, uuid.UUID(job["id"]))
    wb = load_workbook(d / "report.xlsx", read_only=True)
    assert wb.sheetnames == ["Сводка", "Компьютеры", "Уязвимости", "Программы", "Обнаружения", "Параметры отчёта"]       # сводка → компактный список компьютеров → подробности → методика
    rows = list(wb["Обнаружения"].iter_rows(values_only=True))
    assert len(rows) - 1 == 5 and rows[0][:3] == ("Опасность", "Уязвимость (KLA)", "CVE")
    hosts = list(wb["Компьютеры"].iter_rows(values_only=True))
    assert len(hosts) - 1 == 4 and len({r[0] for r in hosts[1:]}) == 4                                                                    # каждый компьютер — одна строка
    summary = [r for r in wb["Сводка"].iter_rows(values_only=True) if r]
    assert any(r[0] == "Затронуто компьютеров (уникальных)" and r[1] == 4 for r in summary) and any(r and r[0] == "Методика, источники и ограничения" for r in wb["Параметры отчёта"].iter_rows(values_only=True))
    html = (d / "report.html").read_text(encoding="utf-8")
    assert html.startswith("<!doctype html>") and "KLA10001" in html and "CVE-2024-0001" in html and "Затронуто компьютеров (уникальных)" in html and "@media print" in html
    assert not re.search(r"<link |<script|src=\"http|href=\"http(?!s://threats)|@import|url\(http", html)                                    # никаких внешних ресурсов и скриптов
    cs = list(csv.reader(io.StringIO((d / "report.csv").read_text(encoding="utf-8-sig")), delimiter=";"))
    assert len(cs) - 1 == 5 and cs[0][0] == "Опасность"
    r = await client.get(f"/api/v1/reports/{job['id']}/download/html")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/html") and "attachment" in r.headers["content-disposition"]
    page = (await client.get(f"/api/v1/reports/{job['id']}/rows", params={"limit": 3, "sort": "name"})).json()
    assert page["total"] == 5 and len(page["items"]) == 3


async def test_hosts_report(client, app):
    sid, runner = await _setup(client, app)
    params = {"host_ids": [6, 1]}
    pv = await preview(client, "quick_hosts", params)
    assert m(pv, "hosts") == 2 and pv["total"] == 2 and {r["name"] for r in pv["items"]} == {"ws-001", "ksc-server"}
    job = (await client.post("/api/v1/reports", json={"kind": "quick_hosts", "title": "Сервер", "params": params, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    wb = load_workbook(job_dir(app.state.settings, uuid.UUID(job["id"])) / "report.xlsx", read_only=True)
    assert {"Компьютеры", "Установленное ПО", "Уязвимости"} <= set(wb.sheetnames)
    sw = list(wb["Установленное ПО"].iter_rows(values_only=True))
    assert len(sw) > 3 and {r[0] for r in sw[1:]} == {"ws-001", "ksc-server"}
    v = list(wb["Уязвимости"].iter_rows(values_only=True))
    assert {r[5] for r in v[1:]} <= {"ws-001", "ksc-server"} and len(v) - 1 == m(pv, "vrows")


async def test_protection_report_matches_page(client, app):
    sid, runner = await _setup(client, app)
    page = (await client.get("/api/v1/protection", params={"problem": "no_av", "online": "false", "limit": 100})).json()
    pv = await preview(client, "quick_protection", {"problem": "no_av"})
    assert pv["total"] == page["total"] == m(pv, "hosts") and {r["host_id"] for r in pv["items"]} == {h["host_id"] for h in page["items"]}
    online = await preview(client, "quick_protection", {"problem": "no_av", "online_only": True})
    page_online = (await client.get("/api/v1/protection", params={"problem": "no_av", "online": "true", "limit": 100})).json()
    assert online["total"] == page_online["total"] and all(r["online"] == "в сети" for r in online["items"])
    job = (await client.post("/api/v1/reports", json={"kind": "quick_protection", "title": "", "params": {"problem": "no_av"}, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    assert (await client.get(f"/api/v1/reports/{job['id']}")).json()["rows"] == page["total"]


async def test_events_report_matches_page(client, app):
    sid, runner = await _setup(client, app)
    page = (await client.get("/api/v1/events", params={"days": 7, "type": "notcured,detected,suspicious", "limit": 50})).json()
    pv = await preview(client, "quick_events", {"days": 7, "types": ["notcured", "detected", "suspicious"]})
    assert pv["total"] == page["total"] and m(pv, "rows") == page["total"] and m(pv, "hosts") >= 1
    one = await preview(client, "quick_events", {"days": 7, "event_ids": [page["items"][0]["event_id"]]})
    assert one["total"] == 1
    job = (await client.post("/api/v1/reports", json={"kind": "quick_events", "title": "", "params": {"days": 7, "group": "notcured"}, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    d = (await client.get(f"/api/v1/reports/{job['id']}")).json()
    assert d["status"] == "done" and d["rows"] == 1


async def test_events_deleted_device_is_labelled(client, app):
    await _setup(client, app)
    pv = await preview(client, "quick_events", {"days": 7})
    gone = [r for r in pv["items"] if r["host_id"] == 9901]
    assert len(gone) == 1 and "ID 9901" in gone[0]["host"] and "имя неизвестно" in gone[0]["host"] and gone[0]["ip"] == "нет данных"
    assert any("которых уже нет в KSC: 1" in n for n in pv["doc"]["notes"]) and "из них устройств, которых уже нет в KSC: 1" in next(m["hint"] for m in pv["doc"]["metrics"] if m["key"] == "hosts")


async def test_html_truncation_is_announced_and_summary_is_full(client, app, monkeypatch):
    monkeypatch.setattr("app.reports.quick_export.HTML_ROW_LIMIT", 2)
    sid, runner = await _setup(client, app)
    params = {"descr_ids": [5001, 5903]}                    # 5 обнаружений, 4 компьютера
    job = (await client.post("/api/v1/reports", json={"kind": "quick_vulns", "title": "", "params": params, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    d = job_dir(app.state.settings, uuid.UUID(job["id"]))
    html = (d / "report.html").read_text(encoding="utf-8")
    assert "сокращённая выборка" in html and "первые 2 из 5 строк" in html and "полному</b> результату" in html
    assert html.count('class="mono">KLA') + html.count('KLA10001</td>') >= 1 and html.index("сокращённая выборка") < html.index("Затронуто компьютеров")
    assert any(x["value"] == 4 for x in (await client.get(f"/api/v1/reports/{job['id']}")).json()["summary"]["metrics"])      # показатели — по полному результату
    wb = load_workbook(d / "report.xlsx", read_only=True)
    assert len(list(wb["Обнаружения"].iter_rows(values_only=True))) - 1 == 5               # в Excel — все строки


async def test_hosts_report_by_page_filter_without_selection_and_constructor_options(client, app):
    """Отчёт по компьютерам можно открыть с условиями страницы (без отметки строк): берутся все подходящие; краткий вариант — только показатели; столбцы выбираются."""
    await _setup(client, app)
    full = await preview(client, "quick_hosts", {"q": "ws-", "include_inactive": True})
    assert full["total"] >= 2 and all(r["name"].startswith("ws-") for r in full["items"]) and m(full, "hosts") == full["total"]
    assert [t["id"] for t in full["doc"]["tables"]] == ["hosts"]                       # по условиям — только состояние защиты; ПО и уязвимости включаются явно (до 500 компьютеров)
    wide = await preview(client, "quick_hosts", {"q": "ws-", "include_inactive": True, "sections": ["software", "protection"]})
    assert "software" in [t["id"] for t in wide["doc"]["tables"]]
    short = await preview(client, "quick_hosts", {"q": "ws-", "include_inactive": True, "detail": "summary"})
    assert [t["id"] for t in short["doc"]["tables"]] == ["summary"] and m(short, "hosts") == full["total"]
    cols = await preview(client, "quick_hosts", {"q": "ws-", "include_inactive": True, "columns": ["name", "ip", "nonexistent"]})
    assert [c["key"] for c in cols["columns"]] == ["name", "ip"]
    none = await client.post("/api/v1/quick/preview", json={"kind": "quick_hosts", "params": {"q": "zzzz-not-found"}})
    assert none.status_code in (404, 409, 422), none.text


async def test_vulns_report_by_page_filter_without_selection(client, app):
    """Отчёт по уязвимостям открывается и без отметки строк: берутся все действующие обнаружения с учётом условий страницы (опасность, поиск по компьютеру)."""
    await _setup(client, app)
    allv = await preview(client, "quick_vulns", {"include_inactive": True})
    crit = await preview(client, "quick_vulns", {"include_inactive": True, "severity": "critical"})
    assert allv["total"] >= crit["total"] > 0 and {r["severity"] for r in crit["items"]} == {"critical"}
    by_host = await preview(client, "quick_vulns", {"include_inactive": True, "q": "ws-001"})
    assert 0 < by_host["total"] <= allv["total"] and {r["name"] for r in by_host["items"]} == {"ws-001"}
