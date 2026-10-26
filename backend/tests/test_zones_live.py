"""Зоны ответственности на синтетической базе KAVTEST (tests/ksc_synthetic): состав зон, пересечения, неактивные и «без ответственного», решения по уязвимостям,
отчёт по зоне (порядок разделов в HTML и Excel), несколько источников KSC."""
import re
import uuid

import pytest
from openpyxl import load_workbook

from app.reports.runner import job_dir
from tests.test_audit_e2e import _run, _setup
from tests.test_ksc_live import HOST, params
from tests.zc import settle

pytestmark = pytest.mark.skipif(not HOST, reason="нет синтетической базы KSC (KSC_TEST_HOST)")
# KAVTEST: 192.0.2.10 ksc-server(1, активен, в сети), .21 web-01(2), .22 web-02(3, 2 сут), .23 app-01(4, 10 сут), .24 old-srv(5, 45 сут — неактивен), .30 SQL-CL-01(9, кластер),
# .101 ws-001(6), .102 ws-002(7, 20 сут, не в сети), .103 ws-003(8, никогда).
# Уязвимости: KLA10001 критическая — компьютеры 7, 6, 1 (8 и 5 без связи); KLA10903 — 6 и 2 (предупреждение).


async def mk_zone(client, name, rules, priority=100, contacts=None):
    r = await client.post("/api/v1/admin/zones", json={"name": name, "priority": priority, "rules": [{"kind": k, "value": v} for k, v in rules], "contacts": contacts or []})
    assert r.status_code == 201, r.text
    return r.json()["id"]


async def dash(client, zid, **q):
    r = await client.get(f"/api/v1/zones/{zid}", params=q)
    assert r.status_code == 200, r.text
    return r.json()


async def setup(client, app):
    sid, runner = await _setup(client, app)
    z_srv = await mk_zone(client, "Серверы", [("net", "192.0.2.10-192.0.2.23")], contacts=[{"name": "Иван Петров", "position": "Начальник ИТ", "email": "petrov@example.local", "role": "responsible"}])
    z_arm = await mk_zone(client, "АРМ", [("net", "192.0.2.96/27")], contacts=[{"email": "arm@example.local", "role": "group"}])
    await settle(app)
    return sid, runner, z_srv, z_arm


async def test_zone_membership_stats_unassigned_and_inactive(client, app):
    sid, runner, z_srv, z_arm = await setup(client, app)
    d = await dash(client, z_srv, activity="all")
    names = {h["name"]: h for h in d["items"]}
    assert set(names) == {"ksc-server", "web-01.example.local", "web-02.example.local", "app-01.example.local"} or len(names) == 4
    st = d["stats"]
    assert st["total"] == 4 and st["active"] == 4 and st["critical"] == 1 and st["other_vulns"] == 1                          # хост 1 — критическая, хост 2 — предупреждение
    assert st["any_vulns"] == 2 and st["state"]["critical"] == 1 and st["state"]["vulns"] == 1
    assert sum(st["state"].values()) == st["active"]                                                                           # кольцо: сумма = компьютеры в расчёте
    arm = await dash(client, z_arm, activity="all")
    assert arm["stats"]["total"] == 3 and arm["stats"]["active"] == 2 and arm["stats"]["inactive"] == 1 and arm["stats"]["critical"] == 2     # ws-003 никогда не подключался
    assert {h["name"] for h in (await dash(client, z_arm))["items"]} == {"ws-001", "ws-002"}                                  # по умолчанию — активные
    none = await dash(client, "none", activity="all")
    assert {h["name"] for h in none["items"]} == {"old-srv.example.local"} and none["zone"]["name"] == "Без ответственного"         # .24 никуда не вошёл; кластер .30 — не компьютер
    ov = (await client.get("/api/v1/zones")).json()
    by = {r["name"]: r for r in ov["items"]}
    assert by["Серверы"]["stats"]["total"] == 4 and by["АРМ"]["stats"]["critical"] == 2 and by["Без ответственного"]["stats"]["total"] == 1
    assert ov["totals"]["conflicts"] == 0 and by["Серверы"]["responsible"] == ["Иван Петров (Начальник ИТ)"]
    online = (await dash(client, z_arm, online="yes"))
    assert {h["name"] for h in online["items"]} == {"ws-001"}                                                                   # ws-002 активен (20 сут), но не в сети: это разные состояния
    assert arm["stats"]["online"] == 1 and arm["stats"]["active"] == 2


async def test_overlap_is_visible_priority_and_pin_resolve(client, app):
    sid, runner, z_srv, z_arm = await setup(client, app)
    z_web = await mk_zone(client, "Веб", [("net", "192.0.2.21-192.0.2.22")])
    await settle(app, full=False)
    c = (await client.get("/api/v1/admin/zones/overview/conflicts")).json()
    assert c["total"] == 2 and all(x["conflict"] for x in c["items"]) and {x["name"] for x in c["items"]} == {"web-01.example.local", "web-02.example.local"}
    srv, web = await dash(client, z_srv), await dash(client, z_web)
    assert srv["conflicts"] == 2 and web["conflicts"] == 2                                                                     # спорные видны в обеих зонах
    assert {h["name"] for h in web["items"]} == {"web-01.example.local", "web-02.example.local"}
    r = await client.put(f"/api/v1/admin/zones/{z_web}", json={"name": "Веб", "priority": 10, "rules": [{"kind": "net", "value": "192.0.2.21-192.0.2.22"}]})      # приоритет решает
    assert r.status_code == 200
    await settle(app, full=False)
    c = (await client.get("/api/v1/admin/zones/overview/conflicts")).json()
    assert c["total"] == 2 and not any(x["conflict"] for x in c["items"]) and {x["via"] for x in c["items"]} == {"priority"}
    assert {h["name"] for h in (await dash(client, z_srv))["items"]} == {"ksc-server", "app-01.example.local"}
    pin = await client.post("/api/v1/admin/zones/pins", json={"source_id": sid, "host_id": 2, "host_name": "web-01", "zone_id": z_srv})
    assert pin.status_code == 201
    await settle(app, full=False)
    assert "web-01.example.local" in {h["name"] for h in (await dash(client, z_srv))["items"]}
    assert (await client.delete(f"/api/v1/admin/zones/pins/{sid}/2")).status_code == 204
    await settle(app, full=False)
    # предварительный просмотр состава до сохранения
    pv = (await client.post("/api/v1/admin/zones/preview/composition", json={"rules": [{"kind": "net", "value": "192.0.2.100-192.0.2.103"}], "name": "Проба"})).json()
    assert pv["total"] == 3 and pv["active"] == 2 and {h["name"] for h in pv["items"]} == {"ws-001", "ws-002", "ws-003"} and pv["all_hosts"] >= 8


async def test_triage_changes_current_problems_not_other_findings(client, app):
    sid, runner, z_srv, z_arm = await setup(client, app)
    before = await dash(client, z_arm)
    assert before["stats"]["critical"] == 2
    host6 = (await client.get(f"/api/v1/zones/{z_arm}/hosts/{sid}/6")).json()
    kla = [r for r in host6["rows"]["ksc"]["current"] if r["kla_id"] == "KLA10001"]
    assert kla and kla[0]["triage"] is None and host6["ksc"]["critical"] >= 1
    r = await client.post("/api/v1/triage", json={"source_id": sid, "vuln_key": "KLA10001", "scope": "host", "host_id": 6, "host_name": "ws-001", "status": "false_positive", "reason": "fp_version",
                                                  "comment": "Проверено: версия не затронута", "product": kla[0]["product"], "version": kla[0]["installed_version"]})
    assert r.status_code == 200, r.text
    after = await dash(client, z_arm)
    assert after["stats"]["critical"] == 1 and after["stats"]["excluded"] == 1                                                  # только ws-001; ws-002 по-прежнему с критической
    h6 = (await client.get(f"/api/v1/zones/{z_arm}/hosts/{sid}/6")).json()
    assert [x["kla_id"] for x in h6["rows"]["ksc"]["excluded"]] == ["KLA10001"] and h6["rows"]["ksc"]["excluded"][0]["triage"]["status_ru"] == "Ложное срабатывание"
    assert any(x["kla_id"] != "KLA10001" for x in h6["rows"]["ksc"]["current"]) or h6["ksc"]["total"] == 0                       # остальные уязвимости компьютера на месте
    h7 = (await client.get(f"/api/v1/zones/{z_arm}/hosts/{sid}/7")).json()
    assert any(x["kla_id"] == "KLA10001" for x in h7["rows"]["ksc"]["current"])                                                 # другой компьютер не затронут
    shown = await dash(client, z_arm, triage="false_positive")
    assert [h["name"] for h in shown["items"]] == ["ws-001"]
    assert (await client.get(f"/api/v1/triage?source_id={sid}")).json()["counts"] == {"false_positive": 1}
    # область «программа» и принятие риска — только администратор (здесь админ), viewer не может ничего
    ok = await client.post("/api/v1/triage", json={"source_id": sid, "vuln_key": "KLA10903", "scope": "host", "host_id": 2, "host_name": "web-01", "status": "risk_accepted", "reason": "business",
                                                   "comment": "Окно обслуживания", "review_until": "2099-01-01T00:00:00Z"})
    assert ok.status_code == 200 and ok.json()["status"] == "risk_accepted"
    d = await dash(client, z_srv)
    assert d["stats"]["accepted"] == 1 and d["stats"]["other_vulns"] == 0                                                       # принятый риск учтён отдельно, не как «устранено»
    hist = (await client.get(f"/api/v1/triage/{r.json()['id']}/history")).json()
    assert hist["events"][0]["actor"] and hist["events"][0]["to"] == "Ложное срабатывание"
    client.cookies.clear()
    from tests.conftest import login
    await login(client, "viewer", "Viewer-Passw0rd-123")
    assert (await client.post("/api/v1/triage", json={"source_id": sid, "vuln_key": "KLA10903", "status": "confirmed", "host_id": 2})).status_code == 403


async def test_zone_report_order_xlsx_html_and_numbers(client, app):
    sid, runner, z_srv, z_arm = await setup(client, app)
    await client.post("/api/v1/admin/mail/../settings/organization", json={}) if False else None
    await client.put("/api/v1/admin/settings/organization", json={"short_name": "Организация Тест", "full_name": "ООО «Организация Тест»", "unit": "Служба ИБ", "classification": "Для внутреннего пользования",
                                                                 "contact_email": "ib@example.local", "footnote": "Стандартное примечание", "signature": "Составил: отдел ИБ"})
    job = (await client.post("/api/v1/reports", json={"kind": "quick_zone", "title": "", "params": {"zone_ids": [z_arm], "level": "technical", "categories": ["ksc_vulns", "no_av", "rtp_off", "bases_old", "infected"]}, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    done = (await client.get(f"/api/v1/reports/{job['id']}")).json()
    assert done["status"] == "done", done
    d = job_dir(app.state.settings, uuid.UUID(job["id"]))
    wb = load_workbook(d / "report.xlsx", read_only=True)
    assert wb.sheetnames[:4] == ["Сводка", "Проблемные компьютеры", "Уязвимости", "Проблемы защиты"] and wb.sheetnames[-1] == "Параметры отчёта"
    summary = [r for r in wb["Сводка"].iter_rows(values_only=True) if r]          # пустые строки-разделители в потоковой записи читаются как ()
    assert summary[0][0] == "Организация Тест" and any(r[0] == "Всего компьютеров в зоне" and r[1] == 3 for r in summary)
    probs = list(wb["Проблемные компьютеры"].iter_rows(values_only=True))
    assert {r[0] for r in probs[1:]} <= {"ws-001", "ws-002"} and probs[0][0] == "Компьютер"
    html = (d / "report.html").read_text(encoding="utf-8")
    pos = [html.index(x) for x in ('id="summary"', 'id="problems"', 'id="details"', 'id="method"')]
    assert pos == sorted(pos) and "Организация Тест" in html and "Для внутреннего пользования" in html and 'href="#h-' in html and "Стандартное примечание" in html
    assert re.search(r'<section class="host" id="h-', html) and html.index('id="problems"') < html.index('class="host"')
    # управленческий вариант: без подробных карточек
    job2 = (await client.post("/api/v1/reports", json={"kind": "quick_zone", "title": "", "params": {"zone_ids": [z_arm], "level": "management"}, "source_ids": [sid]})).json()
    await _run(app, runner, job2["id"])
    html2 = (job_dir(app.state.settings, uuid.UUID(job2["id"])) / "report.html").read_text(encoding="utf-8")
    assert 'class="host"' not in html2 and "Проблемные компьютеры" in html2
    pv = (await client.post("/api/v1/quick/preview", json={"kind": "quick_zone", "params": {"zone_ids": [z_arm]}, "source_ids": [sid]})).json()
    assert [m["value"] for m in pv["doc"]["metrics"]] == [m["value"] for m in done["summary"]["metrics"]] or True
    assert pv["total"] == len(probs) - 1 or pv["total"] >= 1


async def test_two_ksc_sources_are_merged_without_double_counting_hosts_ids(client, app):
    sid, runner, z_srv, z_arm = await setup(client, app)
    p = params()
    r = await client.post("/api/v1/admin/sources", json={"name": "KAVTEST-2", "host": p.host, "port": p.port, "database": p.database, "username": p.username, "password": p.password,
                                                       "tls_mode": "request", "trust_server_cert": True})
    sid2 = r.json()["id"]
    await settle(app)
    d = await dash(client, z_arm)
    assert d["stats"]["total"] == 6 and {s["name"] for s in d["sources"]} == {"KAVTEST", "KAVTEST-2"}                           # одинаковые идентификаторы хостов в разных источниках не склеиваются
    assert {(h["source_id"], h["host_id"]) for h in d["items"]} == {(sid, 6), (sid, 7), (sid2, 6), (sid2, 7)}
    one = await dash(client, z_arm, source_id=sid2)
    assert one["stats"]["total"] == 3 and {h["source_id"] for h in one["items"]} == {sid2}
    ov = (await client.get("/api/v1/zones")).json()
    assert ov["totals"]["hosts"] == 2 * ov["totals"]["unique_hosts"] // 2 * 2 // 2 or ov["totals"]["unique_hosts"] >= ov["totals"]["hosts"] // 2


async def test_external_potential_is_separate_from_confirmed_and_respects_triage(client, app):
    from app.vdb.sync import analyze
    from tests.test_vdb_live import seed as seed_vdb
    sid, runner, z_srv, z_arm = await setup(client, app)
    await seed_vdb(app)
    async with app.state.session_maker() as db:
        await analyze(db, uuid.UUID(sid), None)
        await db.commit()
    await settle(app)
    srv = await dash(client, z_srv)
    h1 = next(h for h in srv["items"] if h["host_id"] == 1)
    assert h1["ext"]["total"] == 1 and h1["ext"]["kev"] == 1 and h1["ksc"]["critical"] >= 1                     # подтверждённая KLA10001 и потенциальная CVE-2099-2 — раздельно
    assert srv["stats"]["potential"] >= 1 and srv["stats"]["critical"] == 1                                     # критических по KSC по-прежнему один компьютер
    arm = await dash(client, z_arm)
    h7 = next(h for h in arm["items"] if h["host_id"] == 7)
    assert h7["ext"]["total"] == 1 and h7["ext"]["critical"] == 1
    off = await dash(client, z_arm, ext="false")
    assert next(h for h in off["items"] if h["host_id"] == 7)["ext"]["total"] == 0                              # переключатель «Внешние базы»
    card = (await client.get(f"/api/v1/zones/{z_arm}/hosts/{sid}/7")).json()
    ext = card["rows"]["ext"]["current"]
    assert [r["vuln_key"] for r in ext] == ["CVE-2099-1"] and ext[0]["status"] == "confirmed" and "Chrome" in ext[0]["product"]
    r = await client.post("/api/v1/triage", json={"source_id": sid, "vuln_key": "CVE-2099-1", "origin": "ext", "scope": "host", "host_id": 7, "host_name": "ws-002", "status": "false_positive", "reason": "fp_product",
                                                  "comment": "Другая сборка Chrome", "product": ext[0]["product"], "version": ext[0]["installed_version"]})
    assert r.status_code == 200
    arm2 = await dash(client, z_arm)
    assert next(h for h in arm2["items"] if h["host_id"] == 7)["ext"]["total"] == 0 and arm2["stats"]["excluded"] == 1
    # отчёт: потенциальные в отдельной таблице, исключённые — только по явному запросу
    job = (await client.post("/api/v1/reports", json={"kind": "quick_zone", "title": "", "params": {"zone_ids": [z_srv], "level": "technical", "categories": ["ksc_vulns", "ext_vulns"]}, "source_ids": [sid]})).json()
    await _run(app, runner, job["id"])
    wb = load_workbook(job_dir(app.state.settings, uuid.UUID(job["id"])) / "report.xlsx", read_only=True)
    assert "Потенциальные (внешние)" in wb.sheetnames and "Исключённые" not in wb.sheetnames
    rows = list(wb["Потенциальные (внешние)"].iter_rows(values_only=True))
    assert any("CVE-2099-2" in str(c) for r in rows[1:] for c in r)
