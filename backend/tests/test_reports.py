"""Построитель отчётов на синтетических данных: подставная сессия KSC отвечает на шаблоны запросов заранее заданными строками."""
from datetime import datetime, timedelta, timezone

from app.ksc.grouping import PHP_BRANCH_REGEX, PHP_NAME_REGEX, Rule
from app.reports import builder, export

NOW = datetime.now(timezone.utc)
PRODUCTS = [
    {"product_id": 1, "name": "php7.1-cli", "version": "7.1.33-14+0~2020", "publisher": "Debian PHP Maintainers", "platform_mask": 2},
    {"product_id": 2, "name": "php7.4-xdebug", "version": "3.2.1-1", "publisher": "Debian PHP Maintainers", "platform_mask": 2},
    {"product_id": 3, "name": "php-common", "version": "2:70+debian", "publisher": "", "platform_mask": 2},
    {"product_id": 4, "name": "PhpStorm 2023.2", "version": "2023.2", "publisher": "JetBrains", "platform_mask": 1},
    {"product_id": 5, "name": "php8.3-fpm", "version": "8.3.6-1", "publisher": "", "platform_mask": 2},
]
INSTALLS = [(10, 1), (10, 2), (10, 3), (11, 5), (12, 1), (13, 4), (14, 5)]
HOSTS = {10: ("web-01", NOW - timedelta(days=1)), 11: ("web-02", NOW - timedelta(days=2)), 12: ("old-01", NOW - timedelta(days=40)), 13: ("ws-01", NOW), 14: ("cluster-01", NOW)}


class FakeSession:
    def query(self, sql, params=None, *, max_rows=100_000):
        params = params or {}
        if "FROM dbo.hst_inventory_products p WHERE p.wstrDisplayName LIKE" in sql:
            rows = [p for p in PRODUCTS if "php" in p["name"].lower()]
            return rows[params["offset"]:params["offset"] + params["limit"]]
        if "FROM dbo.hst_inventory_products p WHERE p.nId IN" in sql:
            ids = {int(x) for x in sql.split("IN (")[1].rstrip(")").split(",")}
            return [p for p in PRODUCTS if p["product_id"] in ids]
        if "FROM dbo.hst_inventory_hstprds i" in sql:
            ids = {int(x) for x in sql.split("IN (")[1].split(")")[0].split(",")}
            rows = [{"install_id": n, "host_id": h, "product_id": p, "install_date": None} for n, (h, p) in enumerate(INSTALLS, 1) if p in ids]
            return rows[params["offset"]:params["offset"] + params["limit"]]
        if "FROM dbo.v_akpub_host h WHERE h.nId IN" in sql:
            ids = {int(x) for x in sql.split("IN (")[1].rstrip(")").split(",")}
            return [{"host_id": i, "name": HOSTS[i][0], "dns_name": HOSTS[i][0], "dns_domain": "example.local", "ip_int": 3221225984 + i, "os_name": "Linux", "platform_type": 6,
                     "group_id": 5, "last_connected": HOSTS[i][1], "is_slave": 1 if i == 14 else 0} for i in ids if i in HOSTS]
        if "FROM dbo.v_akpub_adm_group g" in sql:
            return [{"group_id": 0, "name": "Управляемые устройства", "parent_id": 2}, {"group_id": 5, "name": "Серверы", "parent_id": 0}]
        if "base-files" in sql:
            return [{"host_id": 10, "name": "base-files", "version": "12.4+deb12u12"}]
        raise AssertionError("неожиданный запрос: " + sql[:80])


def _build(params):
    rule = Rule("php", "PHP", "%php%", PHP_NAME_REGEX, PHP_BRANCH_REGEX)
    return builder.build_software(FakeSession(), params, rule, max_rows=1000, exclude_slaves=True, progress=lambda t: None, cancelled=lambda: False)


def test_php_audit_default_activity():
    rows, summary = _build({"rule": "php", "active_days": 7})
    # web-01: 3 пакета; web-02: 1; old-01 исключён (40 суток), ws-01 — PhpStorm не PHP, cluster-01 — не компьютер
    assert summary["rows"] == 4 and summary["hosts"] == 2 and summary["excluded_inactive_hosts"] == 1 and summary["excluded_not_computers"] == 1
    branches = {b["branch"]: b for b in summary["branches"]}
    assert branches["7.1"]["installs"] == 1 and branches["7.4"]["installs"] == 1 and branches["8.3"]["hosts"] == 1 and branches["(не определена)"]["installs"] == 1
    assert rows[0]["linux_release"] == "base-files 12.4+deb12u12" and rows[0]["ip"] == "192.0.2.10" and rows[0]["group"] == "Управляемые устройства / Серверы"


def test_include_inactive_and_branch_filter():
    rows, summary = _build({"rule": "php", "active_days": 7, "include_inactive": True, "branches": ["7.1"]})
    assert summary["rows"] == 2 and {r["name"] for r in rows} == {"web-01", "old-01"} and summary["excluded_inactive_hosts"] == 0


def test_version_range_and_export(tmp_path):
    rows, summary = _build({"rule": "php", "min_version": "7", "below_version": "8", "include_inactive": True})
    assert {r["app_name"] for r in rows} == {"php7.1-cli", "php7.4-xdebug"}
    export.write_xlsx(tmp_path / "r.xlsx", "software_hosts", rows, summary, {"Отчёт": "test"})
    export.write_csv(tmp_path / "r.csv", rows, builder.SOFTWARE_COLUMNS)
    export.write_jsonl(tmp_path / "rows.jsonl", rows)
    from openpyxl import load_workbook  # noqa: PLC0415

    wb = load_workbook(tmp_path / "r.xlsx")
    assert wb.sheetnames == ["Установки", "Версии", "Компьютеры", "Параметры"]
    assert wb["Установки"].max_row == len(rows) + 1
    assert export.read_jsonl(tmp_path / "rows.jsonl", 1, 1)[0]["app_name"] == rows[1]["app_name"]
    assert (tmp_path / "r.csv").read_text(encoding="utf-8-sig").startswith("Компьютер;FQDN;IP-адрес")
