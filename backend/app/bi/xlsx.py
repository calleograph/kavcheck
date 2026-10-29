"""Безопасное потоковое чтение XLSX выгрузки.

* До открытия проверяется архив: число частей, распакованный объём и степень сжатия (защита от «архивных бомб»).
* openpyxl в режиме read_only (строки читаются потоком), data_only (формулы не вычисляются — берутся сохранённые значения), keep_links=False (внешние ссылки не загружаются),
  keep_vba=False; макросы никогда не выполняются. XML разбирается через defusedxml (внешние сущности запрещены)."""
from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Iterator

from .profile import norm_header

MAX_ENTRIES = 5000
MAX_UNPACKED = 1024 * 1024 * 1024
MAX_RATIO = 200
HEADER_SCAN = 10


class XlsxError(ValueError):
    pass


def check_archive(path: Path) -> list[str]:
    """Проверка контейнера до разбора. Возвращает предупреждения (например, о макросах); при опасном содержимом — XlsxError."""
    if not zipfile.is_zipfile(path):
        raise XlsxError("Файл не является книгой XLSX (повреждён, не докопирован или имеет другой формат)")
    warns = []
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if len(infos) > MAX_ENTRIES:
                raise XlsxError(f"В книге слишком много частей ({len(infos)}): разбор отклонён")
            total = sum(i.file_size for i in infos)
            if total > MAX_UNPACKED:
                raise XlsxError(f"Распакованный объём книги {total // (1024 * 1024)} МБ превышает предел {MAX_UNPACKED // (1024 * 1024)} МБ")
            for i in infos:
                if i.compress_size and i.file_size / i.compress_size > MAX_RATIO and i.file_size > 10 * 1024 * 1024:
                    raise XlsxError(f"Подозрительно высокая степень сжатия части «{i.filename}»: похоже на архивную бомбу, разбор отклонён")
            names = {i.filename for i in infos}
            if "xl/workbook.xml" not in names:
                raise XlsxError("В архиве нет книги Excel (xl/workbook.xml)")
            if any(n.endswith("vbaProject.bin") for n in names):
                warns.append("В книге есть макросы: они не выполняются и не читаются")
            if any(n.startswith("xl/externalLinks/") for n in names):
                warns.append("В книге есть внешние ссылки: они не загружаются, используются сохранённые значения")
            bad = z.testzip()
            if bad:
                raise XlsxError(f"Часть книги «{bad}» повреждена")
    except zipfile.BadZipFile as exc:
        raise XlsxError(f"Архив XLSX повреждён: {exc}") from None
    return warns


def open_book(path: Path):
    import openpyxl  # noqa: PLC0415

    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False, keep_vba=False)
    except Exception as exc:  # noqa: BLE001 — любой сбой разбора — понятная причина отказа, а не падение импорта
        raise XlsxError(f"Книгу не удалось прочитать: {type(exc).__name__}: {str(exc)[:200]}") from None


def pick_sheet(wb, name: str):
    if name:
        if name not in wb.sheetnames:
            raise XlsxError(f"В книге нет листа «{name}» (листы: {', '.join(wb.sheetnames[:10])})")
        ws = wb[name]
    else:
        ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()                       # выгрузки BI часто пишут неверный размер листа: читаем до последней строки
    return ws


def headers(ws, header_row: int = 0) -> tuple[int, list[str]]:
    """(номер строки заголовка, заголовки). 0 — первая строка с двумя и более непустыми текстовыми ячейками среди первых десяти. Повторяющиеся заголовки получают суффикс."""
    for n, row in enumerate(ws.iter_rows(min_row=1, max_row=header_row or HEADER_SCAN, values_only=True), start=1):
        if header_row and n != header_row:
            continue
        cells = [norm_header(v) for v in row]
        if header_row or sum(1 for c in cells if c and not c.replace(".", "").isdigit()) >= 2:
            while cells and not cells[-1]:
                cells.pop()
            out, seen = [], {}
            for i, c in enumerate(cells):
                c = c or f"Столбец {i + 1}"
                seen[c] = seen.get(c, 0) + 1
                out.append(c if seen[c] == 1 else f"{c} ({seen[c]})")
            return n, out
    raise XlsxError("Не найдена строка заголовков: в первых строках листа нет названий столбцов")


def rows(ws, header_n: int, hdr: list[str]) -> Iterator[tuple[int, dict]]:
    w = len(hdr)
    for n, row in enumerate(ws.iter_rows(min_row=header_n + 1, values_only=True), start=header_n + 1):
        vals = list(row[:w]) + [None] * (w - len(row))
        yield n, dict(zip(hdr, vals))
